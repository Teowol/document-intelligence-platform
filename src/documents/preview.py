"""Bounded previews of original files, independent of the indexing pipeline."""

import csv
import io
from itertools import islice

from openpyxl import load_workbook

from .models import Document

TEXT_BYTES = 200_000
TABLE_ROWS = 100
TABLE_COLUMNS = 30


def build_preview(document, sheet=None):
    if document.doc_type == Document.DocType.PDF:
        return {"kind": "pdf"}

    with document.file.open("rb") as source:
        if document.doc_type == Document.DocType.XLSX:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                sheet_names = workbook.sheetnames
                selected = sheet if sheet in sheet_names else sheet_names[0]
                worksheet = workbook[selected]
                rows = [
                    ["" if cell is None else str(cell) for cell in row]
                    for row in islice(
                        worksheet.iter_rows(
                            max_col=min(worksheet.max_column or TABLE_COLUMNS + 1, TABLE_COLUMNS + 1),
                            values_only=True,
                        ),
                        TABLE_ROWS + 1,
                    )
                ]
            finally:
                workbook.close()
            return _table(rows, sheets=sheet_names, selected_sheet=selected)

        content = source.read(TEXT_BYTES + 1)
        truncated = len(content) > TEXT_BYTES
        text = content[:TEXT_BYTES].decode("utf-8-sig", errors="replace")
        if document.doc_type == Document.DocType.CSV:
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = list(islice(csv.reader(io.StringIO(text), dialect), TABLE_ROWS + 1))
            return _table(rows, truncated=truncated)
        return {"kind": "text", "text": text, "truncated": truncated}


def _table(rows, **extra):
    truncated = (
        extra.pop("truncated", False)
        or len(rows) > TABLE_ROWS
        or any(len(row) > TABLE_COLUMNS for row in rows)
    )
    return {
        "kind": "table",
        "rows": [row[:TABLE_COLUMNS] for row in rows[:TABLE_ROWS]],
        "truncated": truncated,
        **extra,
    }
